from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.shared import Pt

from app.core.schemas import ExtractionResult


def write_all_exports(result: ExtractionResult, output_dir: str | Path) -> dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "txt": output_dir / "resultado.txt",
        "md": output_dir / "resultado.md",
        "json": output_dir / "resultado.json",
        "docx": output_dir / "resultado.docx",
    }
    paths["txt"].write_text(result.full_text, encoding="utf-8")
    paths["md"].write_text(to_markdown(result), encoding="utf-8")
    paths["json"].write_text(json.dumps(result.model_dump(), ensure_ascii=False, indent=2), encoding="utf-8")
    write_docx(result, paths["docx"])
    return paths


def to_markdown(result: ExtractionResult) -> str:
    meta = result.metadata
    lines = [
        f"# Extracción de texto — {meta.filename}",
        "",
        f"- Páginas extraídas: {meta.selected_start_page}–{meta.selected_end_page} de {meta.total_pages}",
        f"- Modo: {meta.options.get('mode')}",
        f"- OCR: {meta.options.get('ocr_strategy')} / idioma {meta.options.get('language')}",
        "",
    ]
    if result.warnings:
        lines.append("## Avisos")
        lines.extend(f"- {warning}" for warning in result.warnings)
        lines.append("")
    for page in result.pages:
        lines.append(f"## Página {page.page_number}")
        lines.append("")
        if page.report.warnings:
            lines.extend(f"> Aviso: {warning}" for warning in page.report.warnings)
            lines.append("")
        lines.append(page.text.strip())
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def write_docx(result: ExtractionResult, output_path: str | Path) -> None:
    output_path = Path(output_path)
    doc = Document()
    doc.core_properties.title = f"Extracción de texto — {result.metadata.filename}"
    doc.core_properties.subject = "Texto extraído de PDF"
    doc.add_heading(f"Extracción de texto — {result.metadata.filename}", level=1)
    summary = doc.add_paragraph()
    summary.add_run("Páginas extraídas: ").bold = True
    summary.add_run(f"{result.metadata.selected_start_page}–{result.metadata.selected_end_page} de {result.metadata.total_pages}")
    summary.add_run("\nModo: ").bold = True
    summary.add_run(str(result.metadata.options.get("mode")))
    summary.add_run("\nOCR: ").bold = True
    summary.add_run(f"{result.metadata.options.get('ocr_strategy')} / {result.metadata.options.get('language')}")

    if result.warnings:
        doc.add_heading("Avisos", level=2)
        for warning in result.warnings:
            doc.add_paragraph(warning, style="List Bullet")

    for page in result.pages:
        doc.add_page_break()
        doc.add_heading(f"Página {page.page_number}", level=2)
        if page.report.warnings:
            for warning in page.report.warnings:
                paragraph = doc.add_paragraph(style="Intense Quote")
                paragraph.add_run(f"Aviso: {warning}")
        for paragraph_text in _split_docx_paragraphs(page.text):
            paragraph = doc.add_paragraph()
            for idx, line in enumerate(paragraph_text.splitlines()):
                if idx:
                    paragraph.add_run().add_break()
                paragraph.add_run(line)

    for style_name in ("Normal",):
        try:
            style = doc.styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(10.5)
        except Exception:
            pass
    doc.save(output_path)


def _split_docx_paragraphs(text: str) -> Iterable[str]:
    chunks = [chunk.strip() for chunk in text.split("\n\n")]
    return [chunk for chunk in chunks if chunk]


def create_searchable_pdf(input_pdf: str | Path, output_pdf: str | Path, language: str = "spa+eng") -> tuple[bool, str | None]:
    """Create a searchable PDF using OCRmyPDF if the binary is available."""
    input_pdf = Path(input_pdf)
    output_pdf = Path(output_pdf)
    binary = shutil.which("ocrmypdf")
    if not binary:
        return False, "OCRmyPDF no está instalado; no se generó PDF buscable"

    cmd = [
        binary,
        "--skip-text",
        "--deskew",
        "--rotate-pages",
        "--optimize",
        "1",
        "-l",
        language,
        str(input_pdf),
        str(output_pdf),
    ]
    try:
        completed = subprocess.run(cmd, check=False, capture_output=True, text=True, timeout=60 * 30)
    except subprocess.TimeoutExpired:
        return False, "OCRmyPDF superó el tiempo máximo configurado"
    except Exception as exc:
        return False, f"No se pudo ejecutar OCRmyPDF: {exc}"

    if completed.returncode != 0:
        stderr = (completed.stderr or completed.stdout or "").strip().replace("\n", " ")
        return False, f"OCRmyPDF falló: {stderr[:500]}"
    return True, None
